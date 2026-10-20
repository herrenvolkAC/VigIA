from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "informe_accesos_oracle_IT_20261020.md"
OUTPUT = ROOT / "Informe_accesos_Oracle_VigIA_para_IT_20261020.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=105, bottom=90, end=105):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_font(run, name="Aptos", size=None, bold=None, color="000000"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_inline(paragraph, text, size=9.2, color="202020"):
    pieces = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for piece in pieces:
        if not piece:
            continue
        if piece.startswith("`") and piece.endswith("`"):
            run = paragraph.add_run(piece[1:-1])
            set_font(run, "Consolas", size - 0.4, False, "243447")
        elif piece.startswith("**") and piece.endswith("**"):
            run = paragraph.add_run(piece[2:-2])
            set_font(run, "Aptos", size, True, color)
        else:
            run = paragraph.add_run(piece)
            set_font(run, "Aptos", size, False, color)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(9.2)
    normal.font.color.rgb = RGBColor(32, 32, 32)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    title.font.size = Pt(23)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(12)
    title_p_pr = title.element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    for style_name, size, before, after in (
        ("Heading 1", 15, 14, 6),
        ("Heading 2", 11.5, 10, 4),
    ):
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    label_run = paragraph.add_run("Página ")
    set_font(label_run, size=8, color="666666")
    field_run = paragraph.add_run()
    set_font(field_run, size=8, color="666666")
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    field_run._r.append(fld_char1)
    field_run._r.append(instr_text)
    field_run._r.append(fld_char2)


def add_table(doc, rows):
    headers = rows[0]
    body = rows[1:]
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    widths = [Cm(4.5), Cm(9.3), Cm(8.7), Cm(3.9)] if len(headers) == 4 else None

    header = table.rows[0]
    set_repeat_table_header(header)
    prevent_row_split(header)
    for idx, text in enumerate(headers):
        cell = header.cells[idx]
        if widths:
            cell.width = widths[idx]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "243447")
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(text)
        set_font(run, size=8.5, bold=True, color="FFFFFF")

    for ridx, values in enumerate(body):
        row = table.add_row()
        prevent_row_split(row)
        for idx, text in enumerate(values):
            cell = row.cells[idx]
            if widths:
                cell.width = widths[idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if ridx % 2:
                set_cell_shading(cell, "F3F6F8")
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            add_inline(paragraph, text, 7.7)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_object_grid(doc, values):
    values = [value for value in values if value]
    columns = 3
    row_count = (len(values) + columns - 1) // columns
    table = doc.add_table(rows=row_count, cols=columns)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.style = "Table Grid"
    widths = [Cm(8.6)] * columns
    for ridx, row in enumerate(table.rows):
        prevent_row_split(row)
        for cidx, cell in enumerate(row.cells):
            cell.width = widths[cidx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, top=65, start=95, bottom=65, end=95)
            if ridx % 2:
                set_cell_shading(cell, "F3F6F8")
            idx = ridx * columns + cidx
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            if idx < len(values):
                run = paragraph.add_run(values[idx])
                set_font(run, "Consolas", 8.0, False, "243447")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def parse_markdown(doc, text):
    lines = text.splitlines()
    i = 0
    in_code = False
    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        if line.startswith("```"):
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                if lines[i].strip():
                    code_lines.append(lines[i].strip())
                i += 1
            add_object_grid(doc, code_lines)
            i += 1
            continue
        if not line:
            i += 1
            continue
        if line.startswith("# "):
            paragraph = doc.add_paragraph(style="Title")
            paragraph.add_run(line[2:].strip())
            i += 1
            continue
        if line.startswith("## "):
            doc.add_paragraph(line[3:].strip(), style="Heading 1")
            i += 1
            continue
        if line.startswith("### "):
            doc.add_paragraph(line[4:].strip(), style="Heading 2")
            i += 1
            continue
        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:|\-]+\|$", lines[i + 1]):
            table_rows = []
            table_rows.append([cell.strip() for cell in line.strip("|").split("|")])
            i += 2
            while i < len(lines) and lines[i].startswith("|"):
                table_rows.append([cell.strip() for cell in lines[i].strip("|").split("|")])
                i += 1
            add_table(doc, table_rows)
            continue
        if re.match(r"^\d+\.\s+", line):
            paragraph = doc.add_paragraph(style="List Number")
            add_inline(paragraph, re.sub(r"^\d+\.\s+", "", line), 9.0)
            i += 1
            continue
        if line.startswith("- "):
            paragraph = doc.add_paragraph(style="List Bullet")
            add_inline(paragraph, line[2:], 9.0)
            i += 1
            continue
        paragraph = doc.add_paragraph()
        add_inline(paragraph, line, 9.2)
        i += 1


def main():
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Cm(1.55)
    section.bottom_margin = Cm(1.45)
    section.left_margin = Cm(1.65)
    section.right_margin = Cm(1.65)
    configure_styles(doc)

    parse_markdown(doc, SOURCE.read_text(encoding="utf-8"))

    core = doc.core_properties
    core.title = "VigIA Accesos Oracle para un nuevo usuario de aplicación"
    core.subject = "Matriz de objetos Oracle, modalidad y frecuencia para IT"
    core.author = "Ingeniería"
    core.keywords = "VigIA, Oracle, permisos, usuario de aplicación, IT"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
