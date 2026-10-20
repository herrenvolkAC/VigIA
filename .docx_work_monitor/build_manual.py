from pathlib import Path
from copy import deepcopy

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(r"C:\Ingenieria\VigIA")
REFERENCE = ROOT / "Docs" / "Manual_Usuario_Recepcion_Refrigerados.docx"
OUTPUT = ROOT / "Docs" / "Manual_Usuario_Monitor_Cargas.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths):
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)
            set_cell_margins(row.cells[idx])
            row.cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(int(sum(widths) * 1440)))
    tbl_w.set(qn("w:type"), "dxa")
    grid = tbl.tblGrid
    if grid is not None:
        for child in list(grid):
            grid.remove(child)
        for width in widths:
            col = OxmlElement("w:gridCol")
            col.set(qn("w:w"), str(int(width * 1440)))
            grid.append(col)


def mark_header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def style_run(run, size=None, bold=None, color=None, font="Aptos"):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_body(doc, text, style="Normal", bold_prefix=None):
    p = doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        style_run(r, bold=True, color="0E1620")
        r2 = p.add_run(text[len(bold_prefix):])
        style_run(r2)
    else:
        r = p.add_run(text)
        style_run(r)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(text)
    style_run(r)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    r = p.add_run(text)
    style_run(r)
    return p


def add_callout(doc, label, text, fill="E5F3E9"):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    mark_header_row(table.rows[0])
    set_table_widths(table, [6.8])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(label + "\n")
    style_run(r, size=10, bold=True, color="1F7A4D", font="Aptos Display")
    r2 = p.add_run(text)
    style_run(r2, size=10)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_table(doc, headers, rows, widths, header_fill="0B5FAE"):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    mark_header_row(table.rows[0])
    set_table_widths(table, widths)
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, header_fill)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(header)
        style_run(r, size=9, bold=True, color="FFFFFF", font="Aptos Display")
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_margins(cells[idx])
            p = cells[idx].paragraphs[0]
            r = p.add_run(str(value))
            style_run(r, size=9)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def clear_body(doc):
    body = doc._element.body
    sect_pr = body.sectPr
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)


doc = Document(REFERENCE)
clear_body(doc)
section = doc.sections[0]
section.left_margin = Inches(0.85)
section.right_margin = Inches(0.85)
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.70)

doc.core_properties.title = "Manual de Usuario - Monitor de Cargas"
doc.core_properties.subject = "Seguimiento de bultos por dársena"
doc.core_properties.author = "VigIA"

header = section.header.paragraphs[0]
header.text = ""
hr = header.add_run("VigIA · OPERACIÓN")
style_run(hr, size=9, bold=True, color="1F7A4D", font="Aptos Display")
footer = section.footer.paragraphs[0]
footer.text = ""
fr = footer.add_run("Manual de usuario · Monitor de Cargas · Página ")
style_run(fr, size=8, color="5C6773")
page_field = OxmlElement("w:fldSimple")
page_field.set(qn("w:instr"), "PAGE")
page_run = OxmlElement("w:r")
page_rpr = OxmlElement("w:rPr")
page_fonts = OxmlElement("w:rFonts")
page_fonts.set(qn("w:ascii"), "Aptos")
page_fonts.set(qn("w:hAnsi"), "Aptos")
page_rpr.append(page_fonts)
page_run.append(page_rpr)
page_text = OxmlElement("w:t")
page_text.text = "1"
page_run.append(page_text)
page_field.append(page_run)
footer._p.append(page_field)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(2)
r = p.add_run("GUÍA OPERATIVA")
style_run(r, size=9, bold=True, color="0B5FAE", font="Aptos Display")

p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
r = p.add_run("Monitor de Cargas")
style_run(r, size=27, bold=True, color="0E1620", font="Aptos Display")

p = doc.add_paragraph()
r = p.add_run("Manual de usuario para seguimiento de bultos por dársena antes del cierre del camión")
style_run(r, size=12, color="5C6773")

add_callout(doc, "Objetivo", "Visualizar en línea qué se está cargando en cada dársena y comparar los bultos efectivos, los bultos planificados estimados y el target operativo antes de cerrar el camión.")

add_body(doc, "1. Acceso al módulo", "Heading 1")
add_body(doc, "Ingresar al sitio interno de VigIA con el usuario habilitado. Desde el selector de aplicaciones, elegir Monitor Cargas.")
add_body(doc, "El módulo muestra los viajes abiertos en preparación y organiza la información por dársena. El acceso requiere una sesión válida de VigIA.")

add_body(doc, "2. Vista general", "Heading 1")
add_body(doc, "Al ingresar, el monitor consulta la información disponible y presenta un resumen superior seguido de las tarjetas de las dársenas.")
add_table(doc, ["Elemento", "Qué muestra"], [
    ("Bultos cargados", "Suma de los bultos efectivos registrados en los viajes visibles."),
    ("Bultos objetivo", "Referencia operativa calculada para los viajes visibles."),
    ("Cumplimiento de bultos", "Semáforo general según bultos efectivos sobre el objetivo."),
    ("Tarjetas de dársena", "Una tarjeta por viaje/dársena, ordenada para la lectura operativa."),
], [1.8, 5.0])

add_body(doc, "3. Interpretar una tarjeta", "Heading 1")
add_body(doc, "Cada tarjeta reúne la información necesaria para decidir si la carga está avanzando correctamente.")
add_bullet(doc, "Dársena: identificación principal del lugar de carga, mostrada en tamaño grande.")
add_bullet(doc, "Viaje y HR: referencias del viaje operativo y de la hoja de ruta.")
add_bullet(doc, "Bultos efectivos: cantidad de bultos registrados hasta el momento.")
add_bullet(doc, "Bultos planificados: estimación de bultos asociada a la hoja de ruta en el sistema de planificación.")
add_bullet(doc, "Bultos target: referencia operativa utilizada para el semáforo de bultos.")
add_bullet(doc, "Semáforo: indica el nivel de cumplimiento del objetivo de bultos.")
add_bullet(doc, "Estado y calidad: informa si hay carga registrada y si existen observaciones de consistencia.")

add_body(doc, "4. Semáforo de ocupación", "Heading 1")
add_body(doc, "El semáforo se calcula con bultos, no con volumetría ni con cantidad de pallets:")
add_table(doc, ["Color", "Regla", "Lectura operativa"], [
    ("Verde", "Más de 90% del target", "Carga próxima o por encima del objetivo."),
    ("Amarillo", "Entre 60% y 90% inclusive", "Carga intermedia; continuar siguiendo el avance."),
    ("Rojo", "Menos de 60% del target", "Carga baja respecto del objetivo."),
    ("Gris", "Sin target disponible", "No hay referencia suficiente para calcular el semáforo."),
], [1.1, 2.0, 3.7], header_fill="1F7A4D")
add_callout(doc, "Importante", "El semáforo es una señal de seguimiento. Antes del cierre, validar también que la carga física y la documentación del viaje correspondan a la operación.", fill="FBE7C8")

add_body(doc, "5. Filtros y actualización", "Heading 1")
add_body(doc, "La pantalla permite adaptar la lectura al momento operativo:")
add_bullet(doc, "Buscar viaje: escribir una dársena, viaje, hoja de ruta, matrícula o cargador.")
add_bullet(doc, "División: seleccionar una división o dejar Todas.")
add_bullet(doc, "Mostrar todas las cargas: por defecto está desmarcado y muestra sólo viajes con bultos efectivos. Marcarlo para incluir también viajes sin carga iniciada.")
add_bullet(doc, "Refresh automático: por defecto está marcado y consulta cada 1 minuto.")
add_bullet(doc, "Minutos: indicar el intervalo deseado entre 1 y 120 minutos. Al modificarlo, el intervalo se reprograma.")
add_bullet(doc, "Actualizar: ejecutar una consulta manual inmediata cuando se necesite una lectura puntual.")

add_body(doc, "6. Ver el detalle de un viaje", "Heading 1")
add_body(doc, "Presionar o seleccionar una tarjeta para expandir el detalle de sus expediciones.")
add_number(doc, "Ubicar la dársena y verificar viaje, HR y matrícula.")
add_number(doc, "Comparar bultos efectivos, planificados y target.")
add_number(doc, "Revisar el semáforo y el estado de calidad.")
add_number(doc, "Abrir la tarjeta para consultar expediciones, bultos y pallets asociados.")
add_number(doc, "Cerrar el detalle volviendo a seleccionar la tarjeta.")

add_body(doc, "7. Estado “Observación”", "Heading 1")
add_body(doc, "Observación no significa necesariamente que la carga esté mal. Indica que el sistema detectó una diferencia o una inconsistencia en los datos de origen.")
add_bullet(doc, "Diferencia de bultos contra la cabecera de la expedición o del pallet.")
add_bullet(doc, "Pallet cargado sin detalle asociado.")
add_bullet(doc, "Líneas sin variable logística o con nivel logístico inválido.")
add_bullet(doc, "Variables logísticas duplicadas.")
add_body(doc, "Abrir la tarjeta para ver el motivo específico. Si la observación afecta la decisión de cierre, solicitar la revisión del dato en el sistema de origen antes de cerrar el camión.")

add_body(doc, "8. Recomendaciones de uso", "Heading 1")
add_table(doc, ["Situación", "Recomendación"], [
    ("Inicio de la jornada", "Dejar el filtro en Cargas iniciadas para concentrarse en los viajes que ya tienen bultos efectivos."),
    ("Seguimiento", "Usar Refresh automático cada 1 minuto y observar cambios en efectivos, planificados y target."),
    ("Revisión puntual", "Usar Actualizar después de una acción de carga o cuando se necesite confirmar el último dato."),
    ("Sin target", "No interpretar el color gris como carga baja: primero verificar la referencia disponible para el viaje."),
    ("Observación", "Leer el detalle de calidad y escalar sólo cuando la diferencia pueda afectar el cierre."),
], [1.7, 5.1])

doc.save(OUTPUT)
print(OUTPUT)
