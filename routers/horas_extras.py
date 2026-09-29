"""Modulo independiente de aprobacion de horas extras."""
from __future__ import annotations

import io
import json
import os
from datetime import date, datetime
from pathlib import Path
from typing import Any, Literal

import aiosqlite
from fastapi import APIRouter, HTTPException, Query, Request
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field

from db.auth import AUTH_DB_PATH
from db.paths import ROOT_DIR, resolve_db_path
from db.schema import DB_PATH as OPERATIVA_DB_PATH
from routers.auth_local import current_auth

router = APIRouter(prefix="/api/horas-extras", tags=["horas-extras"])

HORAS_EXTRAS_DB_PATH = resolve_db_path("VIGIA_HORAS_EXTRAS_DB_PATH", "horas_extras.db", ROOT_DIR)
SENSITIVE_TOKENS = tuple(
    token.strip().upper()
    for token in os.getenv("VIGIA_HE_SENSITIVE_TOKENS", "F.C,FC,M4").split(",")
    if token.strip()
)
HE_MIN_HOURS = float(os.getenv("VIGIA_HE_MIN_HOURS", "0.5"))

SCHEMA = """
CREATE TABLE IF NOT EXISTS he_eventos (
    evento_id INTEGER PRIMARY KEY AUTOINCREMENT,
    legajo TEXT NOT NULL,
    fecha TEXT NOT NULL,
    fuente_hash TEXT NOT NULL,
    snapshot_json TEXT NOT NULL,
    estado_actual TEXT NOT NULL DEFAULT 'detectada',
    horas_detectadas REAL NOT NULL DEFAULT 0,
    horas_autorizadas REAL NOT NULL DEFAULT 0,
    cargado_sap INTEGER NOT NULL DEFAULT 0,
    sap_usuario TEXT,
    sap_at TEXT,
    sap_observacion TEXT,
    created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
    updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
    UNIQUE(legajo, fecha)
);
CREATE TABLE IF NOT EXISTS he_aprobaciones (
    aprobacion_id INTEGER PRIMARY KEY AUTOINCREMENT,
    evento_id INTEGER NOT NULL REFERENCES he_eventos(evento_id) ON DELETE CASCADE,
    etapa TEXT NOT NULL,
    estado TEXT NOT NULL,
    horas_autorizadas REAL,
    observacion TEXT NOT NULL DEFAULT '',
    usuario TEXT NOT NULL,
    created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
);
CREATE TABLE IF NOT EXISTS he_historial (
    historial_id INTEGER PRIMARY KEY AUTOINCREMENT,
    evento_id INTEGER,
    legajo TEXT NOT NULL,
    fecha TEXT NOT NULL,
    accion TEXT NOT NULL,
    estado_desde TEXT,
    estado_hasta TEXT,
    detalle_json TEXT NOT NULL DEFAULT '{}',
    usuario TEXT NOT NULL,
    created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
);
CREATE TABLE IF NOT EXISTS he_ajustes (
    ajuste_id INTEGER PRIMARY KEY AUTOINCREMENT,
    evento_id INTEGER NOT NULL REFERENCES he_eventos(evento_id) ON DELETE CASCADE,
    estado TEXT NOT NULL DEFAULT 'solicitado',
    horas_solicitadas REAL,
    motivo TEXT NOT NULL DEFAULT '',
    usuario_solicita TEXT NOT NULL,
    usuario_resuelve TEXT,
    resuelto_at TEXT,
    created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
);
CREATE TABLE IF NOT EXISTS he_responsables_sector (
    responsable_id INTEGER PRIMARY KEY AUTOINCREMENT,
    sector TEXT NOT NULL,
    jefe_username TEXT,
    gerente_username TEXT,
    activo INTEGER NOT NULL DEFAULT 1,
    vigencia_desde TEXT,
    vigencia_hasta TEXT,
    updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
    UNIQUE(sector, jefe_username, gerente_username)
);
CREATE INDEX IF NOT EXISTS idx_he_eventos_fecha ON he_eventos(fecha);
CREATE INDEX IF NOT EXISTS idx_he_eventos_estado ON he_eventos(estado_actual);
CREATE INDEX IF NOT EXISTS idx_he_historial_evento ON he_historial(evento_id, created_at);
"""


class DecisionRequest(BaseModel):
    legajo: str = Field(min_length=1, max_length=40)
    fecha: date
    etapa: Literal["pre_autorizacion", "gerencia", "ajuste", "sap"]
    estado: Literal[
        "detectada",
        "pre_autorizada",
        "rechazada_jefe",
        "aprobada_gerente",
        "rechazada_gerente",
        "ajuste_solicitado",
        "ajuste_aprobado",
        "cargada_sap",
    ]
    horas_autorizadas: float | None = Field(default=None, ge=0)
    observacion: str = Field(default="", max_length=800)


class BulkDecisionRequest(BaseModel):
    rows: list[DecisionRequest] = Field(min_length=1, max_length=500)


class ResponsableRequest(BaseModel):
    sector: str = Field(min_length=1, max_length=160)
    jefe_username: str | None = Field(default="", max_length=80)
    gerente_username: str | None = Field(default="", max_length=80)
    activo: bool = True
    vigencia_desde: date | None = None
    vigencia_hasta: date | None = None


async def init_horas_extras_db() -> None:
    HORAS_EXTRAS_DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    async with aiosqlite.connect(HORAS_EXTRAS_DB_PATH) as db:
        await db.execute("PRAGMA busy_timeout = 10000")
        await db.execute("PRAGMA journal_mode = WAL")
        await db.executescript(SCHEMA)
        await db.commit()


async def _require_auth(request: Request) -> dict[str, Any]:
    auth = await current_auth(request)
    if not auth or auth.get("device_status") != "approved":
        raise HTTPException(401, "No autenticado.")
    return auth


async def _access(auth: dict[str, Any]) -> dict[str, Any]:
    if auth.get("role") in {"admin", "rrhh"}:
        return {"profile": "RRHH", "scope": "global", "sector": "", "sectors": []}
    async with aiosqlite.connect(AUTH_DB_PATH) as db:
        db.row_factory = aiosqlite.Row
        async with db.execute(
            """
            SELECT profile, scope, sector, metadata_json
            FROM auth_user_app_access
            WHERE username=? AND module='horas_extras' AND enabled=1
            """,
            (auth["username"],),
        ) as cur:
            row = await cur.fetchone()
    if not row:
        return {"profile": "", "scope": "", "sector": "", "sectors": []}
    data = dict(row)
    try:
        metadata = json.loads(data.get("metadata_json") or "{}")
    except json.JSONDecodeError:
        metadata = {}
    data["sectors"] = metadata.get("sectors") or ([data["sector"]] if data.get("sector") else [])
    data["profile"] = (data.get("profile") or "CONSULTA").upper()
    data["scope"] = (data.get("scope") or "").lower()
    return data


def _can_see_sensitive(auth: dict[str, Any], access: dict[str, Any]) -> bool:
    profile = (access.get("profile") or "").upper()
    return auth.get("role") in {"admin", "rrhh"} or profile in {"RRHH", "GERENTE"}


def _can_admin_config(auth: dict[str, Any], access: dict[str, Any]) -> bool:
    return auth.get("role") in {"admin", "rrhh"} or (access.get("profile") or "").upper() == "RRHH"


async def _legajero_responsables() -> dict[str, list[dict[str, Any]]]:
    personas_query = """
        SELECT
            legajo,
            COALESCE(nombre, '') AS nombre,
            COALESCE(desc_sector_generico, '') AS sector,
            COALESCE(desc_funcion, '') AS funcion,
            COALESCE(desc_posicion, '') AS posicion
        FROM rrhh_personas
        WHERE active = 1
          AND (
            UPPER(COALESCE(desc_funcion, '')) LIKE ?
            OR UPPER(COALESCE(desc_posicion, '')) LIKE ?
          )
        ORDER BY nombre, legajo
    """
    legajero_query = """
        WITH latest AS (
            SELECT l.*,
                   ROW_NUMBER() OVER (
                       PARTITION BY l.legajo
                       ORDER BY l.batch_id DESC, l.id DESC
                   ) rn
            FROM rrhh_legajero l
            WHERE TRIM(COALESCE(l.fecha_baja, '')) = ''
              AND TRIM(COALESCE(l.legajo, '')) <> ''
        )
        SELECT
            legajo,
            COALESCE(nombre, '') AS nombre,
            COALESCE(desc_sector_generico, '') AS sector,
            COALESCE(desc_funcion, '') AS funcion,
            COALESCE(desc_posicion, '') AS posicion
        FROM latest
        WHERE rn = 1
          AND (
            UPPER(COALESCE(desc_funcion, '')) LIKE ?
            OR UPPER(COALESCE(desc_posicion, '')) LIKE ?
          )
        ORDER BY nombre, legajo
    """
    legajero_all_query = """
        WITH latest AS (
            SELECT l.*,
                   ROW_NUMBER() OVER (
                       PARTITION BY l.legajo
                       ORDER BY l.batch_id DESC, l.id DESC
                   ) rn
            FROM rrhh_legajero l
            WHERE TRIM(COALESCE(l.fecha_baja, '')) = ''
              AND TRIM(COALESCE(l.legajo, '')) <> ''
        )
        SELECT
            legajo,
            COALESCE(nombre, '') AS nombre,
            COALESCE(desc_sector_generico, '') AS sector,
            COALESCE(desc_funcion, '') AS funcion,
            COALESCE(desc_posicion, '') AS posicion
        FROM latest
        WHERE rn = 1
        ORDER BY nombre, legajo
    """

    async def fetch_like(db: aiosqlite.Connection, query: str, patterns: tuple[str, str]) -> list[dict[str, Any]]:
        async with db.execute(query, patterns) as cur:
            rows = [dict(row) for row in await cur.fetchall()]
        return unique_people(rows)

    async def fetch_all(db: aiosqlite.Connection, query: str) -> list[dict[str, Any]]:
        async with db.execute(query) as cur:
            rows = [dict(row) for row in await cur.fetchall()]
        return unique_people(rows)

    def unique_people(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
        seen: set[str] = set()
        result: list[dict[str, Any]] = []
        for row in rows:
            legajo = str(row.get("legajo") or "").strip()
            if not legajo or legajo in seen:
                continue
            seen.add(legajo)
            row["legajo"] = legajo
            result.append(row)
        return result

    async with aiosqlite.connect(OPERATIVA_DB_PATH) as db:
        db.row_factory = aiosqlite.Row
        jefes = await fetch_like(db, personas_query, ("%JEFE%", "%JEFE%"))
        gerentes = await fetch_like(db, personas_query, ("%GEREN%", "%GEREN%"))
        if not jefes:
            jefes = await fetch_like(db, legajero_query, ("%JEFE%", "%JEFE%"))
        if not gerentes:
            gerentes = await fetch_like(db, legajero_query, ("%GEREN%", "%GEREN%"))
        if not jefes and not gerentes:
            personas = await fetch_all(db, legajero_all_query)
            jefes = personas
            gerentes = personas
        return {
            "jefes": jefes,
            "gerentes": gerentes,
        }


async def _sector_catalog(auth: dict[str, Any], access: dict[str, Any]) -> dict[str, list[str]]:
    can_sensitive = _can_see_sensitive(auth, access)
    queries = [
        """
        SELECT DISTINCT desc_sector_generico sector, desc_funcion funcion, desc_area_personal
        FROM rrhh_personas
        WHERE active=1 AND TRIM(COALESCE(desc_sector_generico,'')) <> ''
        ORDER BY sector, funcion
        """,
        """
        WITH latest AS (
            SELECT l.*,
                   ROW_NUMBER() OVER (
                       PARTITION BY l.legajo
                       ORDER BY l.batch_id DESC, l.id DESC
                   ) rn
            FROM rrhh_legajero l
            WHERE TRIM(COALESCE(l.fecha_baja, '')) = ''
        )
        SELECT DISTINCT desc_sector_generico sector, desc_funcion funcion, desc_area_personal
        FROM latest
        WHERE rn = 1 AND TRIM(COALESCE(desc_sector_generico,'')) <> ''
        ORDER BY sector, funcion
        """,
        """
        SELECT DISTINCT sector, '' funcion, area_personal desc_area_personal
        FROM rrhh_actividad_diaria
        WHERE TRIM(COALESCE(sector,'')) <> ''
        ORDER BY sector
        """,
    ]
    async with aiosqlite.connect(OPERATIVA_DB_PATH) as db:
        db.row_factory = aiosqlite.Row
        for query in queries:
            async with db.execute(query) as cur:
                rows = [dict(row) for row in await cur.fetchall()]
            sectors: dict[str, list[str]] = {}
            for row in rows:
                if not _sector_allowed(auth, access, row.get("sector")):
                    continue
                if _is_sensitive(row.get("desc_area_personal")) and not can_sensitive:
                    continue
                sector = row.get("sector") or ""
                funcion = row.get("funcion") or ""
                sectors.setdefault(sector, [])
                if funcion and funcion not in sectors[sector]:
                    sectors[sector].append(funcion)
            if sectors:
                return sectors
    return {}


def _is_sensitive(value: str | None) -> bool:
    text = str(value or "").upper()
    return any(token in text for token in SENSITIVE_TOKENS)


def _sector_allowed(auth: dict[str, Any], access: dict[str, Any], sector: str | None) -> bool:
    if auth.get("role") in {"admin", "rrhh"}:
        return True
    scope = access.get("scope") or ""
    if scope == "global":
        return True
    sectors = {s for s in access.get("sectors") or [] if s}
    return not sectors or (sector or "") in sectors


def _normalize_row(row: dict[str, Any], decision: dict[str, Any] | None = None) -> dict[str, Any]:
    horas_detectadas = sum(float(row.get(key) or 0) for key in ("hs_ext_realiz", "hs_50_autorizadas", "hs_100"))
    if horas_detectadas < HE_MIN_HOURS:
        horas_detectadas = 0.0
    estado = (decision or {}).get("estado_actual") or "detectada"
    horas_autorizadas = float((decision or {}).get("horas_autorizadas") or 0)
    if estado == "aprobada_gerente" and not horas_autorizadas:
        horas_autorizadas = horas_detectadas
    return {
        "fecha": row.get("fecha") or "",
        "legajo": row.get("legajo") or "",
        "empleado": row.get("empleado") or row.get("nombre") or "",
        "sector": row.get("sector") or "",
        "funcion": row.get("funcion") or "",
        "desc_area_personal": row.get("desc_area_personal") or "",
        "sensible": _is_sensitive(row.get("desc_area_personal")),
        "ph": row.get("horario_teorico") or row.get("horario") or row.get("regla_plan_horario") or "",
        "inicio_real": row.get("entrada") or "",
        "fin_real": row.get("salida") or "",
        "horas_trabajadas": float(row.get("hs_trab") or 0),
        "horas_extras_detectadas": horas_detectadas,
        "estado": estado,
        "horas_autorizadas": horas_autorizadas,
        "usuario_cambio": (decision or {}).get("usuario_cambio") or "",
        "fecha_cambio": (decision or {}).get("fecha_cambio") or "",
        "observacion": (decision or {}).get("observacion") or "",
        "cargado_sap": bool((decision or {}).get("cargado_sap") or 0),
        "sap_usuario": (decision or {}).get("sap_usuario") or "",
        "sap_at": (decision or {}).get("sap_at") or "",
    }


async def _activity_rows(desde: date, hasta: date, sector: str = "", funciones: list[str] | None = None) -> list[dict[str, Any]]:
    funciones = funciones or []
    where = ["a.fecha BETWEEN ? AND ?"]
    params: list[Any] = [desde.isoformat(), hasta.isoformat()]
    if sector:
        where.append("COALESCE(p.desc_sector_generico, l.desc_sector_generico, a.sector, '') = ?")
        params.append(sector)
    if funciones:
        where.append("COALESCE(p.desc_funcion, l.desc_funcion, '') IN (" + ",".join("?" for _ in funciones) + ")")
        params.extend(funciones)
    async with aiosqlite.connect(OPERATIVA_DB_PATH) as db:
        db.row_factory = aiosqlite.Row
        async with db.execute(
            f"""
            WITH latest_legajero AS (
                SELECT l.*
                FROM rrhh_legajero l
                JOIN (
                    SELECT legajo, MAX(batch_id) batch_id
                    FROM rrhh_legajero
                    GROUP BY legajo
                ) ult ON ult.legajo = l.legajo AND ult.batch_id = l.batch_id
            ),
            actividad_unica AS (
                SELECT *
                FROM (
                    SELECT a.*,
                           ROW_NUMBER() OVER (
                               PARTITION BY a.legajo, a.fecha
                               ORDER BY a.batch_id DESC, a.id DESC
                           ) rn
                    FROM rrhh_actividad_diaria a
                    WHERE a.fecha BETWEEN ? AND ?
                )
                WHERE rn = 1
            )
            SELECT a.legajo, a.fecha, a.empleado,
                   COALESCE(p.desc_sector_generico, l.desc_sector_generico, a.sector, '') sector,
                   COALESCE(p.desc_funcion, l.desc_funcion, '') funcion,
                   COALESCE(p.desc_area_personal, l.desc_area_personal, '') desc_area_personal,
                   COALESCE(p.regla_plan_horario, '') regla_plan_horario,
                   a.horario, a.horario_teorico, a.entrada, a.salida, a.hs_trab,
                   a.hs_ext_realiz, a.hs_50_autorizadas, a.hs_100
            FROM actividad_unica a
            LEFT JOIN rrhh_personas p ON p.legajo = a.legajo
            LEFT JOIN latest_legajero l ON l.legajo = a.legajo
            WHERE {" AND ".join(where[1:]) if len(where) > 1 else "1=1"}
              AND (COALESCE(a.hs_ext_realiz,0)+COALESCE(a.hs_50_autorizadas,0)+COALESCE(a.hs_100,0)) >= ?
            ORDER BY a.fecha DESC, a.legajo
            """,
            tuple(params + [HE_MIN_HOURS]),
        ) as cur:
            return [dict(row) for row in await cur.fetchall()]


async def _decisions(keys: list[tuple[str, str]]) -> dict[tuple[str, str], dict[str, Any]]:
    if not keys:
        return {}
    await init_horas_extras_db()
    placeholders = ",".join("(?,?)" for _ in keys)
    params = [value for key in keys for value in key]
    async with aiosqlite.connect(HORAS_EXTRAS_DB_PATH) as db:
        db.row_factory = aiosqlite.Row
        async with db.execute(
            f"""
            SELECT e.legajo, e.fecha, e.estado_actual, e.horas_autorizadas, e.cargado_sap,
                   e.sap_usuario, e.sap_at,
                   last.usuario usuario_cambio, last.created_at fecha_cambio, last.observacion
            FROM he_eventos e
            LEFT JOIN (
                SELECT a.*
                FROM he_aprobaciones a
                JOIN (
                    SELECT evento_id, MAX(aprobacion_id) aprobacion_id
                    FROM he_aprobaciones
                    GROUP BY evento_id
                ) ult ON ult.aprobacion_id = a.aprobacion_id
            ) last ON last.evento_id = e.evento_id
            WHERE (e.legajo, e.fecha) IN ({placeholders})
            """,
            tuple(params),
        ) as cur:
            rows = [dict(row) for row in await cur.fetchall()]
    return {(row["legajo"], row["fecha"]): row for row in rows}


@router.get("/meta")
async def meta(request: Request):
    auth = await _require_auth(request)
    access = await _access(auth)
    return {
        "user": {"username": auth["username"], "display_name": auth.get("display_name"), "role": auth.get("role")},
        "access": access,
        "can_see_sensitive": _can_see_sensitive(auth, access),
        "sensitive_tokens": SENSITIVE_TOKENS,
    }


@router.get("/filtros")
async def filtros(request: Request):
    auth = await _require_auth(request)
    access = await _access(auth)
    return {"sectores": await _sector_catalog(auth, access)}


@router.get("/config/responsables")
async def listar_responsables(request: Request):
    auth = await _require_auth(request)
    access = await _access(auth)
    if not _can_admin_config(auth, access):
        raise HTTPException(403, "Requiere perfil RRHH o administrador.")
    await init_horas_extras_db()
    async with aiosqlite.connect(HORAS_EXTRAS_DB_PATH) as db:
        db.row_factory = aiosqlite.Row
        async with db.execute(
            """
            SELECT responsable_id, sector, jefe_username, gerente_username, activo,
                   vigencia_desde, vigencia_hasta, updated_at
            FROM he_responsables_sector
            ORDER BY sector, activo DESC, updated_at DESC
            """
        ) as cur:
            rows = [dict(row) for row in await cur.fetchall()]
    sectores = await _sector_catalog(auth, access)
    responsables = await _legajero_responsables()
    return {"rows": rows, "sectores": sorted(sectores.keys()), **responsables}


@router.post("/config/responsables")
async def guardar_responsable(request: Request, payload: ResponsableRequest):
    auth = await _require_auth(request)
    access = await _access(auth)
    if not _can_admin_config(auth, access):
        raise HTTPException(403, "Requiere perfil RRHH o administrador.")
    await init_horas_extras_db()
    async with aiosqlite.connect(HORAS_EXTRAS_DB_PATH) as db:
        await db.execute(
            """
            INSERT INTO he_responsables_sector
                (sector, jefe_username, gerente_username, activo, vigencia_desde, vigencia_hasta, updated_at)
            VALUES (?, ?, ?, ?, ?, ?, CURRENT_TIMESTAMP)
            ON CONFLICT(sector, jefe_username, gerente_username) DO UPDATE SET
                activo=excluded.activo,
                vigencia_desde=excluded.vigencia_desde,
                vigencia_hasta=excluded.vigencia_hasta,
                updated_at=CURRENT_TIMESTAMP
            """,
            (
                payload.sector.strip(),
                (payload.jefe_username or "").strip() or None,
                (payload.gerente_username or "").strip() or None,
                1 if payload.activo else 0,
                payload.vigencia_desde.isoformat() if payload.vigencia_desde else None,
                payload.vigencia_hasta.isoformat() if payload.vigencia_hasta else None,
            ),
        )
        await db.commit()
    return {"ok": True}


@router.delete("/config/responsables/{responsable_id}")
async def borrar_responsable(responsable_id: int, request: Request):
    auth = await _require_auth(request)
    access = await _access(auth)
    if not _can_admin_config(auth, access):
        raise HTTPException(403, "Requiere perfil RRHH o administrador.")
    await init_horas_extras_db()
    async with aiosqlite.connect(HORAS_EXTRAS_DB_PATH) as db:
        await db.execute("DELETE FROM he_responsables_sector WHERE responsable_id=?", (responsable_id,))
        await db.commit()
    return {"ok": True}


@router.get("/registros")
async def registros(
    request: Request,
    desde: date,
    hasta: date,
    sector: str = Query(""),
    funcion: list[str] = Query(default=[]),
):
    auth = await _require_auth(request)
    access = await _access(auth)
    if desde > hasta or (hasta - desde).days > 366:
        raise HTTPException(400, "Seleccione un rango valido de hasta 366 dias.")
    if sector and not _sector_allowed(auth, access, sector):
        raise HTTPException(403, "Sector fuera de su alcance.")
    rows = await _activity_rows(desde, hasta, sector, funcion)
    decisions = await _decisions([(row["legajo"], row["fecha"]) for row in rows])
    can_sensitive = _can_see_sensitive(auth, access)
    normalized = []
    for row in rows:
        item = _normalize_row(row, decisions.get((row["legajo"], row["fecha"])))
        if not _sector_allowed(auth, access, item.get("sector")):
            continue
        if item["sensible"] and not can_sensitive:
            continue
        normalized.append(item)
    return {
        "rows": normalized,
        "quality": {
            "actividad_rows": len(rows),
            "message": "" if rows else "No hay registros de Actividad Diaria con horas extras para el rango consultado.",
        },
    }


def _snapshot(row: dict[str, Any]) -> tuple[str, str]:
    payload = {
        key: row.get(key)
        for key in (
            "fecha",
            "legajo",
            "empleado",
            "sector",
            "funcion",
            "desc_area_personal",
            "ph",
            "inicio_real",
            "fin_real",
            "horas_trabajadas",
            "horas_extras_detectadas",
        )
    }
    text = json.dumps(payload, sort_keys=True, ensure_ascii=True)
    return text, str(abs(hash(text)))


async def _find_current_row(request: Request, legajo: str, fecha: str) -> dict[str, Any]:
    data = await registros(request, date.fromisoformat(fecha), date.fromisoformat(fecha), sector="", funcion=[])
    for row in data["rows"]:
        if row["legajo"] == legajo:
            return row
    raise HTTPException(403, "Registro no disponible o fuera de alcance.")


@router.post("/decisiones")
async def decisiones(request: Request, payload: BulkDecisionRequest):
    auth = await _require_auth(request)
    await init_horas_extras_db()
    async with aiosqlite.connect(HORAS_EXTRAS_DB_PATH) as db:
        db.row_factory = aiosqlite.Row
        await db.execute("BEGIN IMMEDIATE")
        try:
            for item in payload.rows:
                current = await _find_current_row(request, item.legajo, item.fecha.isoformat())
                snapshot, source_hash = _snapshot(current)
                horas = current["horas_extras_detectadas"]
                authorized = horas if item.estado in {"pre_autorizada", "aprobada_gerente", "cargada_sap"} else 0
                if item.horas_autorizadas is not None:
                    authorized = min(float(item.horas_autorizadas), horas)
                async with db.execute(
                    "SELECT evento_id, estado_actual FROM he_eventos WHERE legajo=? AND fecha=?",
                    (item.legajo, item.fecha.isoformat()),
                ) as cur:
                    saved = await cur.fetchone()
                previous = saved["estado_actual"] if saved else None
                await db.execute(
                    """
                    INSERT INTO he_eventos
                        (legajo, fecha, fuente_hash, snapshot_json, estado_actual, horas_detectadas,
                         horas_autorizadas, cargado_sap, sap_usuario, sap_at, sap_observacion, updated_at)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, CURRENT_TIMESTAMP)
                    ON CONFLICT(legajo, fecha) DO UPDATE SET
                        fuente_hash=excluded.fuente_hash,
                        snapshot_json=excluded.snapshot_json,
                        estado_actual=excluded.estado_actual,
                        horas_detectadas=excluded.horas_detectadas,
                        horas_autorizadas=excluded.horas_autorizadas,
                        cargado_sap=excluded.cargado_sap,
                        sap_usuario=COALESCE(excluded.sap_usuario, he_eventos.sap_usuario),
                        sap_at=COALESCE(excluded.sap_at, he_eventos.sap_at),
                        sap_observacion=COALESCE(excluded.sap_observacion, he_eventos.sap_observacion),
                        updated_at=CURRENT_TIMESTAMP
                    """,
                    (
                        item.legajo,
                        item.fecha.isoformat(),
                        source_hash,
                        snapshot,
                        item.estado,
                        horas,
                        authorized,
                        1 if item.estado == "cargada_sap" else current["cargado_sap"],
                        auth["username"] if item.estado == "cargada_sap" else None,
                        datetime.now().strftime("%Y-%m-%d %H:%M:%S") if item.estado == "cargada_sap" else None,
                        item.observacion if item.estado == "cargada_sap" else None,
                    ),
                )
                async with db.execute(
                    "SELECT evento_id FROM he_eventos WHERE legajo=? AND fecha=?",
                    (item.legajo, item.fecha.isoformat()),
                ) as cur:
                    evento = await cur.fetchone()
                evento_id = evento["evento_id"]
                await db.execute(
                    """
                    INSERT INTO he_aprobaciones
                        (evento_id, etapa, estado, horas_autorizadas, observacion, usuario)
                    VALUES (?, ?, ?, ?, ?, ?)
                    """,
                    (evento_id, item.etapa, item.estado, authorized, item.observacion.strip(), auth["username"]),
                )
                await db.execute(
                    """
                    INSERT INTO he_historial
                        (evento_id, legajo, fecha, accion, estado_desde, estado_hasta, detalle_json, usuario)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        evento_id,
                        item.legajo,
                        item.fecha.isoformat(),
                        item.etapa,
                        previous,
                        item.estado,
                        json.dumps({"observacion": item.observacion, "horas_autorizadas": authorized}, ensure_ascii=True),
                        auth["username"],
                    ),
                )
            await db.commit()
        except Exception:
            await db.rollback()
            raise
    return {"guardadas": len(payload.rows)}


def _excel(rows: list[dict[str, Any]]) -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Horas extras"
    headers = [
        "Fecha", "Legajo", "Empleado", "Sector", "Funcion", "PH", "Inicio real", "Fin real",
        "Horas trabajadas", "Horas extras", "Estado", "Horas autorizadas", "Usuario cambio",
        "Fecha cambio", "Observacion", "Cargado SAP",
    ]
    ws.append(headers)
    for row in rows:
        ws.append([row.get(key, "") for key in (
            "fecha", "legajo", "empleado", "sector", "funcion", "ph", "inicio_real", "fin_real",
            "horas_trabajadas", "horas_extras_detectadas", "estado", "horas_autorizadas",
            "usuario_cambio", "fecha_cambio", "observacion", "cargado_sap",
        )])
        for cell in ws[ws.max_row]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F7A4D")
    for i, width in enumerate([13, 12, 28, 24, 24, 18, 14, 14, 18, 16, 18, 18, 18, 20, 42, 14], 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


@router.get("/exportar.xlsx")
async def exportar(request: Request, desde: date, hasta: date, sector: str = "", funcion: list[str] = Query(default=[])):
    data = await registros(request, desde, hasta, sector, funcion)
    return StreamingResponse(
        _excel(data["rows"]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="horas_extras_{desde}_{hasta}.xlsx"'},
    )
