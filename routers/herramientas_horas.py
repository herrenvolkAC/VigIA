"""Decisiones independientes sobre la actividad diaria consolidada de RRHH."""
from datetime import date, datetime, timezone, timedelta
from typing import Literal
import json
import io
import re
import unicodedata

import aiosqlite
from fastapi import APIRouter, HTTPException, Query, Request
from pydantic import BaseModel, Field
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from db.schema import DB_PATH
from routers.rrhh_novedades import CONSOLIDATED_CTES, _require_auth, _visibility_sql

router = APIRouter(prefix='/horas-extras', tags=['herramientas'])
DDL = """
CREATE TABLE IF NOT EXISTS herramientas_he_decisiones (
 legajo TEXT NOT NULL, fecha TEXT NOT NULL, estado TEXT NOT NULL,
 observacion TEXT NOT NULL, origen TEXT NOT NULL, usuario TEXT NOT NULL,
 updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP, version INTEGER NOT NULL,
 PRIMARY KEY(legajo, fecha));
CREATE TABLE IF NOT EXISTS herramientas_he_historial (
 id INTEGER PRIMARY KEY, legajo TEXT NOT NULL, fecha TEXT NOT NULL,
 estado TEXT NOT NULL, observacion TEXT NOT NULL, origen TEXT NOT NULL,
 usuario TEXT NOT NULL, version INTEGER NOT NULL,
 created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP);
"""


class Decision(BaseModel):
    legajo: str = Field(min_length=1, max_length=40)
    fecha: date
    estado: Literal['pendiente', 'autorizada', 'rechazada', 'parcial']
    horas_autorizadas: float | None = Field(default=None, ge=0, allow_inf_nan=False)
    observacion: str = Field(default='', max_length=500)
    origen: str
    version: int = Field(ge=0)


class Guardar(BaseModel):
    rows: list[Decision] = Field(min_length=1, max_length=1000)


def source(row):
    return json.dumps([round(float(row[k] or 0), 6) for k in
                       ('hs_ext_realiz', 'hs_50_autorizadas', 'hs_100')])


async def ensure_decisions(db):
    await db.executescript(DDL)
    await db.execute('BEGIN IMMEDIATE')
    try:
        for table in ('herramientas_he_decisiones', 'herramientas_he_historial'):
            async with db.execute(f'PRAGMA table_info({table})') as cur:
                columns = {row[1] for row in await cur.fetchall()}
            if 'horas_autorizadas' not in columns:
                await db.execute(f'ALTER TABLE {table} ADD COLUMN horas_autorizadas REAL')
        await db.commit()
    except Exception:
        await db.rollback()
        raise


async def activity(db, auth, condition, params):
    visibility = _visibility_sql(auth, 'a', 'l.desc_sector_generico')
    async with db.execute(f"""WITH {CONSOLIDATED_CTES}
        SELECT a.legajo, a.fecha, a.empleado, l.desc_sector_generico sector,
               l.desc_funcion funcion, a.hs_ext_realiz, a.hs_50_autorizadas, a.hs_100
        FROM actividad_consolidada a JOIN latest_legajero l ON l.legajo=a.legajo
        WHERE {visibility} AND {condition}
        AND (COALESCE(a.hs_ext_realiz,0)+COALESCE(a.hs_50_autorizadas,0)+COALESCE(a.hs_100,0))>0
        ORDER BY a.fecha DESC, a.legajo""", params) as cur:
        return [dict(r) for r in await cur.fetchall()]


@router.get('/filtros')
async def filtros(request: Request):
    auth = await _require_auth(request)
    visibility = _visibility_sql(auth, 'l', 'l.desc_sector_generico')
    async with aiosqlite.connect(DB_PATH) as db:
        async with db.execute(f"""WITH {CONSOLIDATED_CTES}
            SELECT DISTINCT l.desc_sector_generico, l.desc_funcion
            FROM latest_legajero l WHERE {visibility}
            AND TRIM(COALESCE(l.desc_sector_generico,''))<>''
            ORDER BY 1,2""") as cur:
            pairs = await cur.fetchall()
    sectors = {}
    for sector, funcion in pairs:
        sectors.setdefault(sector, [])
        if funcion and funcion.strip():
            sectors[sector].append(funcion)
    return {'sectores': sectors}


@router.get('')
async def consultar(request: Request, desde: date, hasta: date,
                    sector: str = Query(min_length=1), funcion: list[str] = Query(default=[])):
    auth = await _require_auth(request)
    if desde > hasta or (hasta-desde).days > 365:
        raise HTTPException(400, 'Seleccione un rango válido de hasta 366 días.')
    condition = 'a.fecha BETWEEN ? AND ? AND l.desc_sector_generico=?'
    params = [desde.isoformat(), hasta.isoformat(), sector]
    if funcion:
        condition += ' AND l.desc_funcion IN (' + ','.join('?' for _ in funcion) + ')'
        params.extend(funcion)
    async with aiosqlite.connect(DB_PATH) as db:
        db.row_factory = aiosqlite.Row
        await ensure_decisions(db)
        rows = await activity(db, auth, condition, params)
        async with db.execute('SELECT * FROM herramientas_he_decisiones WHERE fecha BETWEEN ? AND ?',
                              (desde.isoformat(), hasta.isoformat())) as cur:
            decisions = {(r['legajo'], r['fecha']): dict(r) for r in await cur.fetchall()}
    for row in rows:
        row['origen'] = source(row)
        row['horas'] = sum(json.loads(row['origen']))
        saved = decisions.get((row['legajo'], row['fecha']), {})
        row['revision'] = bool(saved and saved['origen'] != row['origen'])
        row['estado'] = 'pendiente' if row['revision'] else saved.get('estado', 'pendiente')
        row['horas_autorizadas'] = (row['horas'] if row['estado'] == 'autorizada' else
                                   float(saved.get('horas_autorizadas') or 0) if row['estado'] == 'parcial' else 0)
        for key, default in [('version', 0), ('observacion', ''), ('usuario', ''), ('updated_at', '')]:
            row[key] = saved.get(key, default)
    return {'rows': rows}


SORT_FIELDS = Literal['fecha', 'legajo', 'empleado', 'sector', 'funcion',
                      'horas', 'horas_autorizadas', 'estado', 'observacion', 'updated_at']


def argentina_time(value):
    if not value:
        return ''
    stamp = datetime.fromisoformat(value)
    if stamp.tzinfo is None:
        stamp = stamp.replace(tzinfo=timezone.utc)
    return stamp.astimezone(timezone(timedelta(hours=-3))).strftime('%d/%m/%Y %H:%M:%S')


def excel_content(rows, orden='fecha', direccion='desc'):
    def key(row):
        if orden in ('horas', 'horas_autorizadas'):
            return row.get(orden, 0)
        text = unicodedata.normalize('NFKD', str(row.get(orden) or '').casefold())
        text = ''.join(c for c in text if not unicodedata.combining(c))
        return [(0, int(s)) if s.isdigit() else (1, s) for s in re.split(r'(\d+)', text)]

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Horas extras'
    sheet.append(['Fecha', 'Legajo', 'Empleado', 'Sector', 'Función', 'HE realizadas',
                  'Decisión', 'Observación', 'Requiere revisión', 'Usuario', 'Última decisión (hora Argentina)', 'Horas autorizadas'])
    for row in sorted(rows, key=key, reverse=direccion == 'desc'):
        values = [date.fromisoformat(row['fecha']), str(row['legajo']), row['empleado'],
                  row['sector'], row['funcion'], row['horas']/24,
                  {'pendiente':'Pendiente', 'autorizada':'Autorizada', 'rechazada':'Rechazada', 'parcial':'Aprobación parcial'}[row['estado']],
                  row['observacion'], 'Sí' if row['revision'] else 'No', row['usuario'], argentina_time(row['updated_at']), row.get('horas_autorizadas', 0)/24]
        sheet.append(values)
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                # Employee names and observations are literal text, never Excel formulas.
                cell.data_type = 's'
            cell.alignment = Alignment(vertical='top', wrap_text=True)
        sheet.cell(sheet.max_row, 1).number_format = 'dd/mm/yyyy'
        sheet.cell(sheet.max_row, 6).number_format = '[h]:mm'
        sheet.cell(sheet.max_row, 12).number_format = '[h]:mm'
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='1F7A4D')
    for i, width in enumerate([14,14,32,24,28,18,22,50,20,24,24,20], 1):
        sheet.column_dimensions[get_column_letter(i)].width = width
    sheet.freeze_panes = 'C2'
    sheet.auto_filter.ref = sheet.dimensions
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


@router.get('/exportar.xlsx')
async def exportar(request: Request, desde: date, hasta: date,
                   sector: str = Query(min_length=1), funcion: list[str] = Query(default=[]),
                   orden: SORT_FIELDS = 'fecha', direccion: Literal['asc', 'desc'] = 'desc'):
    data = await consultar(request, desde, hasta, sector, funcion)
    return StreamingResponse(excel_content(data['rows'], orden, direccion),
        media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': f'attachment; filename="horas_extras_{desde}_{hasta}.xlsx"'})


@router.post('/decisiones')
async def guardar(request: Request, payload: Guardar):
    auth = await _require_auth(request)
    keys = [(r.legajo, r.fecha) for r in payload.rows]
    if len(keys) != len(set(keys)):
        raise HTTPException(400, 'Hay filas duplicadas.')
    async with aiosqlite.connect(DB_PATH) as db:
        db.row_factory = aiosqlite.Row
        await ensure_decisions(db)
        await db.execute('BEGIN IMMEDIATE')
        for row in payload.rows:
            if row.estado == 'rechazada' and not row.observacion.strip():
                raise HTTPException(400, 'Indique el motivo del rechazo.')
            current = await activity(db, auth, 'a.legajo=? AND a.fecha=?',
                                     (row.legajo, row.fecha.isoformat()))
            if not current:
                raise HTTPException(403, 'Fila no disponible o fuera de su alcance.')
            if source(current[0]) != row.origen:
                raise HTTPException(409, 'Las horas cambiaron. Consulte nuevamente antes de guardar.')
            async with db.execute('SELECT version FROM herramientas_he_decisiones WHERE legajo=? AND fecha=?',
                                  (row.legajo, row.fecha.isoformat())) as cur:
                saved = await cur.fetchone()
            if (saved['version'] if saved else 0) != row.version:
                raise HTTPException(409, 'Otro usuario modificó la decisión. Consulte nuevamente.')
            total = sum(json.loads(source(current[0])))
            authorized = total if row.estado == 'autorizada' else 0
            if row.estado == 'parcial':
                if row.horas_autorizadas is None or not 0 < row.horas_autorizadas < total:
                    raise HTTPException(400, 'La aprobación parcial requiere horas mayores a cero y menores a las realizadas.')
                authorized = row.horas_autorizadas
            values = (row.legajo, row.fecha.isoformat(), row.estado, row.observacion.strip(),
                      row.origen, auth['username'], row.version+1, authorized)
            await db.execute('''INSERT INTO herramientas_he_decisiones
                (legajo,fecha,estado,observacion,origen,usuario,version,horas_autorizadas) VALUES (?,?,?,?,?,?,?,?)
                ON CONFLICT(legajo,fecha) DO UPDATE SET estado=excluded.estado,
                observacion=excluded.observacion,origen=excluded.origen,usuario=excluded.usuario,
                horas_autorizadas=excluded.horas_autorizadas,version=excluded.version,updated_at=CURRENT_TIMESTAMP''', values)
            await db.execute('''INSERT INTO herramientas_he_historial
                (legajo,fecha,estado,observacion,origen,usuario,version,horas_autorizadas) VALUES (?,?,?,?,?,?,?,?)''', values)
        await db.commit()
    return {'guardadas': len(payload.rows)}
