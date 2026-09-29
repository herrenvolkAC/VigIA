import sys,asyncio,base64,io
from pathlib import Path
from datetime import datetime as D
from unittest.mock import patch,AsyncMock
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from routers.herramientas import _calculate_route_payment as calc, _consolidate_excel_rows, recalcular_excel,ExcelRecalcularRequest,exportar_excel,ExcelExportRequest
from openpyxl import Workbook,load_workbook
stops=[{'sucursal':'118','llegada_sucursal':'2026-09-07T03:58:18','salida_sucursal':'2026-09-07T04:49:15','km_ida':514,'km_regreso':514}]
start=D.fromisoformat('2026-09-06T20:38:03');end=D.fromisoformat('2026-09-07T16:46:53')
r=calc(stops,start,end,[]);assert r['km_al_100']==257,r
assert r['km_tramos'][0]['porcentaje_pago']==50
multi=[{'sucursal':'1','llegada_sucursal':'2026-09-05T09:00:00','salida_sucursal':'2026-09-05T10:00:00','km_ida':100,'km_regreso':100}, {'sucursal':'2','llegada_sucursal':'2026-09-05T11:00:00','salida_sucursal':'2026-09-05T12:00:00','km_ida':200,'km_regreso':300}]
r=calc(list(reversed(multi)),D(2026,9,5,8),D(2026,9,5,14),[])
assert r['km_al_100']==150,r
assert r['km_tramos'][1]['sucursal']=='2'
incomplete=calc([multi[0],{**multi[1],'llegada_sucursal':'','salida_sucursal':''}],D(2026,9,5,8),D(2026,9,5,14),[])
assert incomplete['km_al_100'] is None and 'sucursal 2' in incomplete['km_revision'],incomplete
assert calc(stops,start,end,[{'fecha':'2026-09-06','motivo':'Test','activo':1}])['km_al_100'] is None
assert calc([{**stops[0],'km_ida':None}],start,end,[])['km_al_100'] is None
saturday_after_13=[
 {'sucursal':'165','llegada_sucursal':'2026-08-30T07:04:00','salida_sucursal':'2026-08-30T07:32:00','km_ida':340,'km_regreso':340},
 {'sucursal':'96','llegada_sucursal':'2026-08-30T08:44:00','salida_sucursal':'2026-08-30T09:11:00','km_ida':340,'km_regreso':340},
]
r=calc(saturday_after_13,D(2026,8,29,21,31),D(2026,8,30,14,7),[])
assert r['km_al_100']==680,r
assert r['km_tramos'][0]['porcentaje_real']==100,r
assert r['km_tramos'][0]['porcentaje_pago']==100,r
assert r['km_tramos'][1]['porcentaje_pago']==100,r
async def test():
 raw=[{'HOJARUTA':'1967463','SUCURSAL':'118','FESALIDA':start,'FEVUELTA':end,'LLEGADA_SUCURSAL':stops[0]['llegada_sucursal'],'SALIDA_SUCURSAL':stops[0]['salida_sucursal'],'KM_IDA':514,'KM_REGRESO':514,'NOMBRE':'Test','NRODOC':'1'}]
 with patch('routers.herramientas._load_holidays',AsyncMock(return_value=[])):
  rows=await _consolidate_excel_rows([{'hojaruta':'1967463','cantidad':1008,'fecha_viaje':'2026-09-06'}],raw)
  assert rows[0]['cantidad']==1008 and rows[0]['km_al_100']==252
  assert rows[0]['km_oracle_al_100']==257 and rows[0]['km_oracle_total']==1028
  assert rows[0]['porcentaje_global_100']==25
  result=await recalcular_excel(ExcelRecalcularRequest(rows=rows));assert result['rows'][0]['km_al_100']==252
  multi_row={'hojaruta':'1965921','cantidad':670,'fesalida':'2026-09-05T08:00:00','fevuelta':'2026-09-05T14:00:00','nombre':'Test','nrodoc':'1','sucursales':multi}
  multi_result=await recalcular_excel(ExcelRecalcularRequest(rows=[multi_row]))
  assert multi_result['rows'][0]['km_al_100']==268,multi_result
  assert multi_result['rows'][0]['demora_max_minutos']==60,multi_result
  full_row={'hojaruta':'1965910','cantidad':670,'fesalida':'2026-08-29T21:31:00','fevuelta':'2026-08-30T14:07:00','nombre':'Test','nrodoc':'1','sucursales':saturday_after_13}
  full_result=await recalcular_excel(ExcelRecalcularRequest(rows=[full_row]))
  assert full_result['rows'][0]['km_al_100']==670,full_result
  assert full_result['rows'][0]['km_oracle_al_100']==680,full_result
  assert full_result['rows'][0]['porcentaje_global_100']==100,full_result
  wb=Workbook();wb.active.append(['Num.Hoj.Ru','Fecha viaj','Cantidad']);wb.active.append(['1967463','2026-09-06',1008]);b=io.BytesIO();wb.save(b)
  response=await exportar_excel(ExcelExportRequest(filename='test.xlsx',content_b64=base64.b64encode(b.getvalue()).decode(),rows=rows))
  content=b''.join([chunk async for chunk in response.body_iterator]);out=load_workbook(io.BytesIO(content));assert out.active['C2'].value==1008;assert out.active['D2'].value==252
 print('OK: Oracle proportion, Excel-liquidated KM, multiple stops, review messages and export preserved')
asyncio.run(test())
