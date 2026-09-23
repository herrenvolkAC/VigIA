from contextlib import closing
import sqlite3
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import AsyncMock, patch
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from fastapi import HTTPException
from routers import herramientas_horas as h
from openpyxl import load_workbook


class HorasTest(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.path = str(Path(self.tmp.name) / 'test.db')
        with closing(sqlite3.connect(self.path)) as db:
            db.executescript('''
                CREATE TABLE rrhh_import_batches(batch_id INTEGER, imported_at TEXT, status TEXT);
                CREATE TABLE rrhh_legajero(id INTEGER, batch_id INTEGER, legajo TEXT,
                  desc_sector_generico TEXT, desc_funcion TEXT);
                CREATE TABLE rrhh_actividad_diaria(id INTEGER, batch_id INTEGER, legajo TEXT,
                  fecha TEXT, empleado TEXT, hs_ext_realiz REAL, hs_50_autorizadas REAL, hs_100 REAL);
                INSERT INTO rrhh_import_batches VALUES(1,'2026-09-01','complete'),(2,'2026-09-02','complete');
                INSERT INTO rrhh_legajero VALUES(1,1,'1','Secos','Armador'),(2,2,'1','Secos','Armador'),(3,2,'2','Frio','Otro');
                INSERT INTO rrhh_actividad_diaria VALUES(1,1,'1','2026-09-01','Uno',9,0,0),(2,2,'1','2026-09-01','Uno',1,2,0.5);
            ''')
        self.patches = [patch.object(h, 'DB_PATH', self.path), patch.object(h, '_require_auth', AsyncMock(return_value={'role':'admin','username':'tester'}))]
        for p in self.patches: p.start()

    async def asyncTearDown(self):
        for p in reversed(self.patches): p.stop()
        self.tmp.cleanup()

    async def query(self):
        return (await h.consultar(None, date(2026,9,1),date(2026,9,2),'Secos',[]))['rows']

    async def test_decisions_and_stale_source(self):
        rows = await self.query()
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['horas'], 3.5)
        decision = h.Decision(**{**rows[0], 'estado':'autorizada'})
        await h.guardar(None, h.Guardar(rows=[decision]))
        self.assertEqual((await self.query())[0]['estado'], 'autorizada')
        with self.assertRaises(HTTPException) as error:
            await h.guardar(None, h.Guardar(rows=[decision]))
        self.assertEqual(error.exception.status_code,409)
        with closing(sqlite3.connect(self.path)) as db:
            self.assertEqual(db.execute('SELECT COUNT(*) FROM herramientas_he_historial').fetchone()[0],1)
            db.execute('UPDATE rrhh_actividad_diaria SET hs_ext_realiz=2 WHERE id=2')
            db.commit()
        row = (await self.query())[0]
        self.assertTrue(row['revision'])
        self.assertEqual(row['estado'],'pendiente')
        with self.assertRaises(HTTPException):
            await h.guardar(None,h.Guardar(rows=[decision]))

    async def test_filters_scope_rejection_and_atomicity(self):
        self.assertEqual((await h.filtros(None))['sectores']['Secos'],['Armador'])
        row = (await self.query())[0]
        with self.assertRaises(HTTPException):
            await h.guardar(None,h.Guardar(rows=[h.Decision(**{**row,'estado':'rechazada'})]))
        with patch.object(h,'_require_auth',AsyncMock(return_value={'role':'user','username':'scoped','rrhh_scope':'sector_completo','rrhh_sectors':['Frio']})):
            self.assertEqual(await self.query(),[])
            with self.assertRaises(HTTPException) as error:
                await h.guardar(None,h.Guardar(rows=[h.Decision(**{**row,'estado':'autorizada'})]))
            self.assertEqual(error.exception.status_code,403)
        valid = h.Decision(**{**row,'estado':'autorizada'})
        invalid = h.Decision(**{**row,'legajo':'missing','estado':'autorizada'})
        with self.assertRaises(HTTPException):
            await h.guardar(None,h.Guardar(rows=[valid,invalid]))
        self.assertEqual((await self.query())[0]['version'],0)

    async def test_partial_validation_persistence_history_export(self):
        row = (await self.query())[0]
        for hours in (None, 0, row['horas'], row['horas']+1):
            with self.assertRaises(HTTPException) as error:
                await h.guardar(None, h.Guardar(rows=[h.Decision(**{**row, 'estado':'parcial', 'horas_autorizadas':hours})]))
            self.assertEqual(error.exception.status_code, 400)
        await h.guardar(None, h.Guardar(rows=[h.Decision(**{**row, 'estado':'parcial', 'horas_autorizadas':1.5})]))
        saved = (await self.query())[0]
        self.assertEqual(saved['estado'], 'parcial')
        self.assertEqual(saved['horas_autorizadas'], 1.5)
        with closing(sqlite3.connect(self.path)) as db:
            self.assertEqual(db.execute('SELECT horas_autorizadas FROM herramientas_he_historial').fetchone()[0],1.5)
        book = load_workbook(h.excel_content([saved]))
        self.assertEqual(book.active['L2'].value.total_seconds(), 5400)
        book.close()
        await h.guardar(None, h.Guardar(rows=[h.Decision(**{**saved,'estado':'autorizada'})]))
        full = (await self.query())[0]
        self.assertEqual(full['horas_autorizadas'],full['horas'])
        with closing(sqlite3.connect(self.path)) as db:
            db.execute('UPDATE rrhh_actividad_diaria SET hs_ext_realiz=4 WHERE id=2')
            db.commit()
        revised = (await self.query())[0]
        self.assertEqual(revised['estado'], 'pendiente')
        self.assertEqual(revised['horas_autorizadas'],0)

    def test_argentina_time(self):
        self.assertEqual(h.argentina_time('2026-09-15 18:28:08'), '15/09/2026 15:28:08')
        self.assertEqual(h.argentina_time('2026-09-15 01:28:08'), '14/09/2026 22:28:08')
        self.assertEqual(h.argentina_time(''), '')

    async def test_bulk_save_and_excel(self):
        with closing(sqlite3.connect(self.path)) as db:
            db.execute("INSERT INTO rrhh_actividad_diaria VALUES(3,2,'1','2026-09-02','Uno',0,0,2)")
            db.commit()
        rows = await self.query()
        result = await h.guardar(None, h.Guardar(rows=[h.Decision(**{
            **r, 'estado':'rechazada', 'observacion':'=HYPERLINK("external")'
        }) for r in rows]))
        self.assertEqual(result['guardadas'], 2)
        rows = await self.query()
        self.assertTrue(all(r['estado']=='rechazada' for r in rows))
        workbook = load_workbook(h.excel_content(rows, 'horas', 'asc'))
        sheet = workbook.active
        self.assertEqual(sheet.max_row, 3)
        self.assertEqual(sheet['F2'].value.total_seconds(), 7200)
        self.assertEqual(sheet['F2'].number_format, '[h]:mm')
        self.assertEqual(sheet['B2'].data_type, 's')
        self.assertEqual(sheet['H2'].data_type, 's')
        self.assertEqual(sheet['H2'].value, '=HYPERLINK("external")')
        self.assertEqual(sheet.auto_filter.ref, 'A1:L3')
        self.assertEqual(sheet.freeze_panes, 'C2')
        workbook.close()
        with patch.object(h,'_require_auth',AsyncMock(return_value={'role':'user','username':'scoped','rrhh_scope':'sector_completo','rrhh_sectors':['Frio']})):
            response = await h.exportar(None,date(2026,9,1),date(2026,9,2),'Secos',[], 'fecha','asc')
            import io
            workbook = load_workbook(io.BytesIO(b''.join([chunk async for chunk in response.body_iterator])))
            self.assertEqual(workbook.active.max_row, 1)
            workbook.close()


if __name__ == '__main__':
    unittest.main()
