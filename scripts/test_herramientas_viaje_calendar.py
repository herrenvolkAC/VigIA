import asyncio, base64, io, sys
from pathlib import Path
from unittest.mock import AsyncMock, patch
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from openpyxl import Workbook, load_workbook
from routers.herramientas import viaje_calendar, exportar_excel, ExcelExportRequest

async def test():
    cases=[('2026-09-14','Lunes','FFFFFF'),('2026-09-15','Martes','FFFFFF'),('2026-09-16','Miércoles','FFFFFF'),('2026-09-17','Jueves','FFFFFF'),('2026-09-18','Viernes','FFFF00'),('2026-09-19','Sábado','FFC000'),('2026-09-20','Domingo','92D050'),('2026-08-17','Lunes','FCE4D6')]
    holidays=[{'fecha':'2026-08-17','activo':1}]
    book=Workbook();ws=book.active;ws.append(['Num.Hoj.Ru','Fecha viaj'])
    rows=[]
    for i,(day,name,color) in enumerate(cases,1):
        assert viaje_calendar(day,holidays)==(name,color)
        ws.append([str(i),day]);rows.append({'hojaruta':str(i),'fecha_viaje':day,'hoja_fisica_presente':i==1})
    assert viaje_calendar('18/09/2026',[])==('Viernes','FFFF00')
    assert viaje_calendar('2026-02-30',[])==('','FFFFFF')
    assert viaje_calendar('2026-09-20',[{'fecha':'2026-09-20','activo':1}])==('Domingo','FCE4D6')
    buffer=io.BytesIO();book.save(buffer)
    with patch('routers.herramientas._load_holidays',AsyncMock(return_value=holidays)):
        response=await exportar_excel(ExcelExportRequest(filename='test.xlsx',content_b64=base64.b64encode(buffer.getvalue()).decode(),rows=rows))
    import re
    assert re.search(r'Viajes_Larga_Distancia_test_validado_\d{4}-\d{2}-\d{2}_\d{2}-\d{2}-\d{2}\.xlsx',response.headers['content-disposition'])
    output=b''.join([chunk async for chunk in response.body_iterator]);result=load_workbook(io.BytesIO(output));sheet=result.active
    assert sheet.cell(1,8).value=='Día Viaje'
    assert sheet.cell(1,9).value=='Hoja Física Presente'
    assert sheet.cell(2,9).value=='Sí'
    assert sheet.cell(3,9).value=='No'
    for i,(_,name,color) in enumerate(cases,2):
        assert sheet.cell(i,8).value==name
        assert all(c.fill.fgColor.rgb[-6:]==color for c in sheet[i])
    print('OK: seven weekdays, holiday priority, invalid dates, Excel values and full-row colors')
asyncio.run(test())
